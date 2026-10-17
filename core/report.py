"""GOST-style DOCX/XLSX report generation for PV-Calc Pro."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "output"

# Formula image key for each of the 14 calculation steps produced by core.calculator.
STEP_FORMULAS: list[str | None] = [
    "w0",               # 1. W_ac = Σ(P_i·t_i)
    "w1_prof",          # 2. W1 = W_ac/η_inv
    "w_day",            # 3. W_day = (W1+W_dc)/7
    "e_sigma",          # 4. E_sigma = W_day/U_sys
    "c_req",            # 5. C_req = W_day·n/(U_sys·DoD·k_t)
    "n_total_prof",     # 6. Конфигурация АКБ
    "p_pv",             # 7. P_pv = W_day/(PSH·η·k)
    "n_pv",             # 8. N_pv = ceil(P_pv/P_panel)
    "voc_string",       # 9. Холодный Voc стринга
    "i_charge",         # 10. I_charge = P_pv/U_sys
    "inverter_surge",   # 11. Инвертор и пусковые токи
    "cable",            # 12. Кабель DC
    "cable",            # 13. Кабель AC
    None,               # 14. Аппараты защиты (без формулы)
]


def _fmt(value: Any, digits: int = 2, default: str = "—") -> str:
    try:
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return default


def _text(value: Any, default: str = "—") -> str:
    if value is None:
        return default
    value = str(value).strip()
    return value if value else default


def _resolve_output(filename: str | Path, suffix: str) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = Path(filename)
    if output_path.suffix.lower() != suffix:
        output_path = output_path.with_suffix(suffix)
    if not output_path.is_absolute():
        output_path = OUTPUT_DIR / output_path.name
    output_path.parent.mkdir(parents=True, exist_ok=True)
    return output_path


def _setup_document() -> docx.document.Document:
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.18)
        section.right_margin = Inches(0.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def _add_heading(doc: docx.document.Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True


def _add_load_table(doc: docx.document.Document, title: str, loads: list[dict[str, Any]]) -> None:
    doc.add_paragraph(title)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Прибор", "P, Вт", "t, ч/нед", "k_пуск", "Энергия, Вт·ч/нед")
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for load in loads:
        row = table.add_row().cells
        power = float(load.get("power", 0.0))
        hours = float(load.get("hours_week", 0.0))
        row[0].text = _text(load.get("name"))
        row[1].text = _fmt(power, 1)
        row[2].text = _fmt(hours, 1)
        row[3].text = _fmt(load.get("surge_factor", 1.0), 1)
        row[4].text = _fmt(power * hours, 2)
    doc.add_paragraph()


def _add_bom_table(doc: docx.document.Document, results: dict[str, Any]) -> None:
    bom_items = list(results.get("bom", []))
    if not bom_items:
        doc.add_paragraph("BOM не сформирован.")
        return
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Элемент", "Модель", "Кол-во", "Цена, BYN", "Сумма, BYN", "Примечание")
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for item in bom_items:
        row = table.add_row().cells
        row[0].text = _text(item.get("element"))
        row[1].text = _text(item.get("model"))
        row[2].text = f"{_fmt(item.get('quantity'), 2)} {_text(item.get('unit'), '')}"
        row[3].text = _fmt(item.get("unit_price_byn"), 2)
        row[4].text = _fmt(item.get("total_byn"), 2)
        row[5].text = _text(item.get("note"), "")
    doc.add_paragraph()


def _add_economics(doc: docx.document.Document, results: dict[str, Any]) -> None:
    economics = dict(results.get("economics", {}))
    if not economics:
        return
    _add_heading(doc, "5. ЭКОНОМИКА И УДЕЛЬНАЯ СТОИМОСТЬ ЭНЕРГИИ")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Показатель"
    table.rows[0].cells[1].text = "Значение"
    rows = (
        ("Стоимость оборудования", f"{_fmt(economics.get('equipment_subtotal_byn'), 2)} BYN"),
        ("CAPEX с монтажом", f"{_fmt(economics.get('capex_byn'), 2)} BYN"),
        ("Годовая выработка/потребление", f"{_fmt(economics.get('annual_energy_kwh'), 2)} кВт·ч/год"),
        ("Срок службы проекта", f"{_fmt(economics.get('lifetime_years'), 0)} лет"),
        ("Тариф сети", f"{_fmt(economics.get('electricity_tariff_byn_kwh'), 4)} BYN/кВт·ч"),
        ("Удельная стоимость энергии за срок службы", f"{_fmt(economics.get('cost_per_kwh_byn'), 4)} BYN/кВт·ч"),
        ("Простой срок окупаемости", f"{_text(economics.get('simple_payback_years'))} лет"),
    )
    for name, value in rows:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = value
    doc.add_paragraph()


def _add_picture_path(doc: docx.document.Document, asset_path: str, width: float) -> None:
    """Insert a rendered formula PNG (asset path) into the document if it exists."""
    path = BASE_DIR / asset_path
    if not path.exists():
        return
    try:
        doc.add_picture(str(path), width=Inches(width))
    except Exception:  # noqa: BLE001
        pass


def _add_formula_picture(
    doc: docx.document.Document,
    formula_images: dict[str, str] | None,
    key: str,
    width: float = 3.2,
) -> None:
    """Insert a rendered formula PNG into the document if available."""
    if not formula_images or key not in formula_images:
        return
    _add_picture_path(doc, formula_images[key], width)


def _add_methodology(
    doc: docx.document.Document,
    detailed: bool,
    formula_images: dict[str, str] | None,
) -> None:
    """Add calculation methodology section with rendered formulas."""

    _add_heading(doc, "3. МЕТОДИКА РАСЧЁТА")

    if not detailed:
        doc.add_paragraph(
            "Академический уровень — строгое следование методичке БГУИР с фиксированными "
            "коэффициентами потерь, DoD = 0.2 и температурной поправкой 1.2."
        )
        items = [
            ("1. Энергопотребление отдельного прибора", "wi", "Затраты энергии одной нагрузкой за время её работы в неделю."),
            ("2. Суммарная энергия переменного тока за неделю", "w0", "Сумма энергий всех приборов переменного тока."),
            ("3. Энергия с учетом потерь в инверторе", "w1", "Коэффициент потерь 1.2 соответствует КПД инвертора около 80 %."),
            ("4. Суточное энергопотребление", "e_sut_acad", "Перевод недельной энергии в суточный эквивалент с учётом шин AC и DC."),
            ("5. Расчетная емкость АКБ", "c_acad", "Ёмкость с учётом дней автономности и глубины разряда 20 %."),
            ("6. Полная емкость АКБ", "c_poln_acad", "Запас ёмкости на температурную деградацию (коэффициент 1.2)."),
            ("7. Количество параллельных АКБ", "n_akb_acad", "Число батарей, необходимое для покрытия расчётной ёмкости."),
            ("8. Количество последовательных АКБ", "n_posl_acad", "Число батарей в серии для согласования с шинным напряжением."),
            ("9. Пиковые солнечные часы", "tp_acad", "Пересчёт месячной декабрьской инсоляции в средние пиковые часы."),
            ("10. Мощность солнечной батареи", "p_sun_acad", "Требуемая пиковая мощность PV с запасом коэффициента 1.2."),
            ("11. Общее количество модулей", "nm_acad", "Отношение требуемой мощности к мощности одного модуля."),
            ("12. Количество модулей в серии", "nm_posl_acad", "Согласование напряжения стринга с напряжением системы."),
            ("13. Ток в цепи нагрузки", "i_load_acad", "Расчётный ток для выбора сечения кабеля."),
            ("14. Сопротивление кабеля", "r_cable_acad", "Предельное погонное сопротивление при допустимом падении напряжения."),
        ]
    else:
        doc.add_paragraph(
            "Детальный инженерный анализ: уровень 1 повторяет академические формулы, "
            "уровень 2 добавляет профессиональные проверки PSH, DoD по химии АКБ, "
            "пусковые токи, холодный Voc и выбор сечения по ПУЭ."
        )
        items = [
            ("Блок 1. Суточный баланс — академический уровень", "e_sut_acad", "Недельная энергия с фиксированным коэффициентом потерь, делённая на 7 дней."),
            ("Блок 1. Суточный баланс — профессиональный уровень", "w_sut_prof", "Точный КПД инвертора и раздельный учёт потерь на шинах AC и DC."),
            ("Блок 2. Емкость АКБ — академический уровень", "c_poln_acad", "Фиксированные глубина разряда 20 % и температурный запас 1.2."),
            ("Блок 2. Емкость АКБ — профессиональный уровень", "c_req_prof", "Глубина разряда выбирается по химии накопителя (LiFePO4 / GEL / AGM)."),
            ("Блок 2. Конфигурация АКБ", "n_total_prof", "Последовательно-параллельная сборка под напряжение системы и требуемую ёмкость."),
            ("Блок 3. Мощность PV — академический уровень", "p_sun_acad", "Расчёт через среднемесячную инсоляцию декабря."),
            ("Блок 3. Мощность PV — профессиональный уровень", "p_pv_prof", "Пиковые солнечные часы и системный КПД: пыль, нагрев, MPPT, провода."),
            ("Блок 3. Количество модулей", "n_pv_prof", "Округление требуемой мощности до целого числа панелей вверх."),
            ("Блок 4. Холодный холостой ход стринга", "voc_prof", "Проверка, что напряжение стринга на морозе не превысит максимальный вход MPPT."),
            ("Блок 4. Пусковые токи инвертора", "p_pusk_prof", "Пиковая мощность с учётом пуска самой тяжёлой индуктивной нагрузки."),
            ("Блок 4. Сечение кабеля по ПУЭ", "s_cable_prof", "Выбор сечения одновременно по падению напряжения и длительно допустимому току."),
        ]

    for title, key, text in items:
        p = doc.add_paragraph()
        p.add_run(title).bold = True
        _add_formula_picture(doc, formula_images, key)
        doc.add_paragraph(text)
    doc.add_paragraph()


def generate_report(
    filename: str | Path,
    input_data: dict[str, Any],
    results: dict[str, Any],
    formula_images: dict[str, str] | None = None,
    step_images: dict[int, str] | None = None,
) -> Path:
    """Generate a DOCX report and return the saved path."""
    output_path = _resolve_output(filename, ".docx")
    doc = _setup_document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(
        "ЛАБОРАТОРНАЯ РАБОТА №1\n"
        "РАСЧЕТ И ФУНКЦИОНАЛЬНЫЙ КОНТРОЛЬ ПАРАМЕТРОВ АВТОНОМНОЙ ФОТОЭЛЕКТРИЧЕСКОЙ СИСТЕМЫ"
    )
    title_run.bold = True
    title_run.font.size = Pt(14)
    doc.add_paragraph()

    goal = doc.add_paragraph()
    goal.add_run("Цель работы: ").bold = True
    goal.add_run(
        "рассчитать автономную фотоэлектрическую систему с аккумуляторным накопителем, "
        "подобрать PV-модули, АКБ, MPPT-контроллер, инвертор, кабели и аппараты защиты."
    )

    _add_heading(doc, "1. ИСХОДНЫЕ ДАННЫЕ")
    params = doc.add_paragraph()
    params.add_run(f"• Локация: {_text(results.get('city'))}, месяц: {_text(results.get('month'))}\n")
    params.add_run(f"• Рельеф/затенение: {_text(results.get('relief_name'))} (k={_fmt(results.get('relief_coefficient'), 2)})\n")
    params.add_run(f"• Напряжение системы: {_fmt(input_data.get('system_voltage'), 0)} В\n")
    params.add_run(f"• Дней автономности: {_fmt(input_data.get('days_autonomy'), 1)}\n")
    params.add_run(f"• Пиковая инсоляция PSH: {_fmt(results.get('psh', input_data.get('psh')), 2)} ч/сутки\n")
    params.add_run(f"• КПД инвертора: {_fmt(results.get('inverter_efficiency'), 2)}\n")
    params.add_run(f"• Системный КПД PV: {_fmt(results.get('pv_system_efficiency'), 2)}\n")
    params.add_run(f"• Модель PV-модуля: {_text(results.get('panel_model'))}\n")
    params.add_run(f"• Модель АКБ: {_text(results.get('battery_model'))}")

    _add_load_table(doc, "Таблица 1.1 — Нагрузки переменного тока", list(input_data.get("ac_loads", [])))
    dc_loads = list(input_data.get("dc_loads", []))
    if dc_loads:
        _add_load_table(doc, "Таблица 1.2 — Нагрузки постоянного тока", dc_loads)

    _add_heading(doc, "2. РЕЗУЛЬТАТЫ ЭНЕРГЕТИЧЕСКОГО РАСЧЕТА")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение"
    protection = dict(results.get("protection", {}))
    rows = (
        ("Суммарная энергия AC", f"{_fmt(results.get('total_w_ac'))} Вт·ч/нед"),
        ("Энергия с учетом КПД инвертора W1", f"{_fmt(results.get('w1_energy'))} Вт·ч/нед"),
        ("Суммарная энергия DC", f"{_fmt(results.get('total_w_dc'))} Вт·ч/нед"),
        ("Суточное энергопотребление", f"{_fmt(results.get('daily_energy_wh'))} Вт·ч/сутки"),
        ("Требуемая емкость АКБ", f"{_fmt(results.get('battery_capacity_total'))} А·ч"),
        ("Фактическая конфигурация АКБ", f"{_fmt(results.get('battery_series_count'), 0)}S×{_fmt(results.get('battery_parallel_count'), 0)}P = {_fmt(results.get('battery_count'), 0)} шт.; {_fmt(results.get('battery_capacity_actual'), 1)} А·ч"),
        ("Требуемая мощность PV", f"{_fmt(results.get('p_panel_req'))} Вт"),
        ("Установленное PV-поле", f"{_fmt(results.get('nm_modules'), 0)} шт.; {_fmt(results.get('pv_series_count'), 0)}S×{_fmt(results.get('pv_parallel_strings'), 0)}P; {_fmt(results.get('pv_actual_w'), 0)} Вт"),
        ("Холодный Voc стринга", f"{_fmt(results.get('pv_cold_string_voc'), 1)} В"),
        ("Контроллер заряда", f"{_text(results.get('controller_model'))}; I={_fmt(results.get('controller_current'), 0)} А"),
        ("Инвертор", f"{_text(results.get('inverter_model'))}; Pnom={_fmt(results.get('inverter_rated_w'), 0)} Вт; Psurge={_fmt(results.get('inverter_surge_w'), 0)} Вт"),
        ("Кабель DC", f"{_fmt(results.get('cable_section_mm2'), 1)} мм²; ΔU={_fmt(results.get('cable_voltage_drop_v'), 3)} В ({_fmt(results.get('cable_voltage_drop_percent'), 2)}%)"),
        ("Кабель AC", f"{_fmt(results.get('ac_cable_section_mm2'), 1)} мм²; ΔU={_fmt(results.get('ac_cable_voltage_drop_v'), 3)} В ({_fmt(results.get('ac_cable_voltage_drop_percent'), 2)}%)"),
        ("Автомат PV-поля", f"{_text(protection.get('pv_array_breaker_a'))} А"),
        ("Автомат контроллер-АКБ", f"{_text(protection.get('controller_battery_breaker_a'))} А"),
        ("Главный автомат АКБ", f"{_text(protection.get('battery_main_breaker_a'))} А"),
        ("Автомат AC-выхода", f"{_text(protection.get('ac_output_breaker_a'))} А; {_text(protection.get('ac_rcd'))}"),
    )
    for name, value in rows:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = value
    doc.add_paragraph()

    _add_methodology(doc, bool(input_data.get("detailed_analysis")), formula_images)

    _add_heading(doc, "4. СПЕЦИФИКАЦИЯ ОБОРУДОВАНИЯ (BOM)")
    _add_bom_table(doc, results)
    _add_economics(doc, results)

    warnings = list(results.get("warnings", []))
    if warnings:
        _add_heading(doc, "6. ИНЖЕНЕРНЫЕ ПРЕДУПРЕЖДЕНИЯ")
        for warning in warnings:
            doc.add_paragraph(f"• {warning}")
        doc.add_paragraph()

    _add_heading(doc, "7. ХОД РАСЧЕТА")
    steps = list(results.get("steps", []))
    for index, step in enumerate(steps):
        description = str(step).split(":", 1)[0]
        description_paragraph = doc.add_paragraph()
        description_paragraph.add_run(description).bold = True
        step_image = step_images.get(index) if step_images else None
        if step_image:
            _add_picture_path(doc, step_image, width=4.8)
        else:
            key = STEP_FORMULAS[index] if index < len(STEP_FORMULAS) else None
            if key:
                _add_formula_picture(doc, formula_images, key, width=3.0)
            doc.add_paragraph(str(step))

    doc.save(output_path)
    return output_path


def export_bom_xlsx(filename: str | Path, results: dict[str, Any]) -> Path:
    """Export BOM and economics to XLSX. Requires openpyxl."""
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "Для экспорта XLSX установите пакет openpyxl: pip install openpyxl"
        ) from exc

    bom = list(results.get("bom", []))
    if not bom:
        raise ValueError("BOM пуст: сначала выполните расчет.")
    economics = dict(results.get("economics", {}))
    output_path = _resolve_output(filename, ".xlsx")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    sheet.append(["Элемент", "Модель", "Кол-во", "Ед.", "Цена, BYN", "Сумма, BYN", "Примечание"])
    for item in bom:
        sheet.append(
            [
                item.get("element"),
                item.get("model"),
                float(item.get("quantity", 0.0)),
                item.get("unit"),
                float(item.get("unit_price_byn", 0.0)),
                float(item.get("total_byn", 0.0)),
                item.get("note", ""),
            ]
        )

    econ_sheet = workbook.create_sheet("Economics")
    econ_sheet.append(["Показатель", "Значение"])
    for key, value in economics.items():
        econ_sheet.append([key, value])

    workbook.save(output_path)
    return output_path
